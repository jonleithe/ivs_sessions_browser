# flake8: noqa
# isort: skip_file
"""
Defined in xlsx_export.py

Filename:       xlsx_export.py
Author:         jole
Created:        20.05.2026

Description:    Holds helpers for exporting visible IVS sessions to XLSX.
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import defs as D
from .tui import Tui


COLOR_TO_HEX = {
    "black"     : "FF000000",
    "blue"      : "FF1C7ED6",
    "cyan"      : "FF0C8599",
    "green"     : "FF2B8A3E",
    "magenta"   : "FFAE3EC9",
    "red"       : "FFE03131",
    "white"     : "FFFFFFFF",
    "yellow"    : "FFB7791F",
}
INVISIBLE_ON_WHITE = {"white"}

SHEET_FONT_NAME = "Courier New"
HEADER_FILL    = PatternFill("solid", fgColor="FFD9EAF7")
HEADER_FONT    = Font(name=SHEET_FONT_NAME, bold=True, color="FF000000")
HEADER_ALIGN   = Alignment(vertical="top", wrap_text=False)
CELL_ALIGN     = Alignment(vertical="top", wrap_text=False)
WEEK_ALIGN     = Alignment(horizontal="center", vertical="center", wrap_text=False)

MEDIUM_SIDE = Side(style="medium", color="FF000000")
THIN_SIDE = Side(style="thin", color="FF000000")

TABLE_START_ROW = 5
TABLE_START_COL = 2

EXPORTED_STATIONS = frozenset({"Nn", "Ns"})
REMOVED_STATIONS_RE = re.compile(r"\[([^]]*)\]")
STATION_TOKEN_RE = re.compile(r"[A-Z][a-z]")
STATION_HEADER = " Stations "
STATION_COLUMN_WIDTH = len(STATION_HEADER)
WEEK_COLUMN_WIDTH = 11
TRANSFER_COLUMN_WIDTH = 14


def _selected_indices(pretty_print: str | list[str]) -> list[int]:
    if pretty_print == "ALL":
        return list(range(len(D.HEADERS)))

    indices = [
        D.FIELD_INDEX[D.PRETTY_PRINT_COLUMN_TO_FIELD[col_name]]
        for col_name in pretty_print
        if col_name in D.PRETTY_PRINT_COLUMN_TO_FIELD
    ]
    return indices or list(range(len(D.HEADERS)))
# ─── END OF _selected_indices() ───────────────────────────────────────────────


def _row_font_color(
    values: list[str],
    operator_bindings: dict[str, str],
    operator_colors: dict[str, str],
) -> str:
    op_label = values[D.FIELD_INDEX.get("op", 0)].strip()
    for op_key, label in operator_bindings.items():
        if label == op_label:
            color_name = operator_colors.get(op_key, "").lower()
            if color_name in INVISIBLE_ON_WHITE:
                return "FF000000"
            return COLOR_TO_HEX.get(color_name, "FF000000")
    return "FF000000"
# ─── END OF _row_font_color() ────────────────────────────────────────────────


def _formula_string(value: str) -> str:
    return value.replace('"', '""')
# ─── END OF _formula_string() ────────────────────────────────────────────────


def _selected_stations(value: str) -> str:
    """Return a fixed-width Nn/Ns display with removed brackets preserved."""

    removed_blocks = REMOVED_STATIONS_RE.findall(value)
    active_source = REMOVED_STATIONS_RE.sub("", value)

    active = set(STATION_TOKEN_RE.findall(active_source)) & EXPORTED_STATIONS
    removed = {
        token
        for block in removed_blocks
        for token in STATION_TOKEN_RE.findall(block)
    } & EXPORTED_STATIONS

    left = "[Nn]" if "Nn" in removed else "Nn" if "Nn" in active else ""
    right = "[Ns]" if "Ns" in removed else "Ns" if "Ns" in active else ""

    if not left:
        return right.rjust(STATION_COLUMN_WIDTH) if right else ""
    if not right:
        return left
    return f"{left}{right:>{STATION_COLUMN_WIDTH - len(left)}}"


def _iso_week(start_value: str) -> tuple[int, int] | None:
    """Return the ISO year/week from the displayed session start value."""

    try:
        iso = datetime.fromisoformat(start_value.strip()).isocalendar()
    except ValueError:
        return None
    return iso.year, iso.week


def _apply_week_frames(
    ws,
    week_keys: list[tuple[int, int] | None],
    last_column: int,
) -> None:
    """Merge and frame consecutive rows belonging to the same ISO week."""

    if not week_keys:
        return

    group_start = TABLE_START_ROW + 1
    for offset in range(1, len(week_keys) + 1):
        at_end = offset == len(week_keys)
        changed = not at_end and week_keys[offset] != week_keys[offset - 1]
        if not at_end and not changed:
            continue

        group_end = TABLE_START_ROW + offset
        if group_end > group_start:
            ws.merge_cells(
                start_row=group_start,
                start_column=1,
                end_row=group_end,
                end_column=1,
            )

        for row_num in range(group_start, group_end + 1):
            week_cell = ws.cell(row=row_num, column=1)
            week_cell.border = Border(
                left=MEDIUM_SIDE,
                right=MEDIUM_SIDE,
                top=MEDIUM_SIDE if row_num == group_start else Side(),
                bottom=MEDIUM_SIDE if row_num == group_end else Side(),
            )

            for col_num in range(TABLE_START_COL, last_column):
                cell = ws.cell(row=row_num, column=col_num)
                cell.border = Border(
                    top=MEDIUM_SIDE if row_num == group_start else Side(),
                    bottom=MEDIUM_SIDE if row_num == group_end else Side(),
                )

            transfer_cell = ws.cell(row=row_num, column=last_column)
            transfer_cell.border = Border(
                left=THIN_SIDE,
                right=MEDIUM_SIDE,
                top=MEDIUM_SIDE if row_num == group_start else THIN_SIDE,
                bottom=MEDIUM_SIDE if row_num == group_end else THIN_SIDE,
            )

        group_start = group_end + 1


def sessions_to_xlsx_bytes(
    rows: list[D.Row],
    pretty_print: str | list[str],
    operator_bindings: dict[str, str],
    operator_colors: dict[str, str],
) -> bytes:
    """
    Convert visible session rows to an XLSX workbook represented as bytes.
    """

    selected_indices = _selected_indices(pretty_print)
    stations_idx = D.FIELD_INDEX.get("stations")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sessions"
    ws.freeze_panes = ws.cell(row=TABLE_START_ROW + 1, column=1).coordinate

    week_header = ws.cell(row=TABLE_START_ROW, column=1, value="Week")
    week_header.fill = HEADER_FILL
    week_header.font = HEADER_FONT
    week_header.alignment = HEADER_ALIGN
    week_header.border = Border(left=MEDIUM_SIDE, top=MEDIUM_SIDE)

    for col_num, idx in enumerate(selected_indices, start=TABLE_START_COL):
        header = STATION_HEADER if idx == stations_idx else D.HEADERS[idx][0]
        cell = ws.cell(row=TABLE_START_ROW, column=col_num, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = Border(top=MEDIUM_SIDE)

    transfer_col = TABLE_START_COL + len(selected_indices)
    transfer_header = ws.cell(row=TABLE_START_ROW, column=transfer_col, value="Transfer done")
    transfer_header.fill = HEADER_FILL
    transfer_header.font = HEADER_FONT
    transfer_header.alignment = HEADER_ALIGN
    transfer_header.border = Border(right=MEDIUM_SIDE, top=MEDIUM_SIDE)

    type_idx = D.FIELD_INDEX.get("type", 1)
    start_idx = D.FIELD_INDEX.get("start", 3)
    display_values: dict[tuple[int, int], str] = {}
    week_keys: list[tuple[int, int] | None] = []
    for row_num, (values, url, meta) in enumerate(rows, start=TABLE_START_ROW + 1):
        font_color = _row_font_color(values, operator_bindings, operator_colors)
        row_font = Font(name=SHEET_FONT_NAME, color=font_color)

        week_key = _iso_week(values[start_idx])
        week_keys.append(week_key)
        week_cell = ws.cell(
            row=row_num,
            column=1,
            value=f"Week {week_key[1]}" if week_key else "",
        )
        week_cell.font = row_font
        week_cell.alignment = WEEK_ALIGN

        for col_num, idx in enumerate(selected_indices, start=TABLE_START_COL):
            value = values[idx].strip()
            marker = Tui.type_marker(meta) if idx == type_idx else ""
            if marker:
                value = f"{value} {marker}".strip()
            if idx == stations_idx:
                value = _selected_stations(value)

            display_values[(row_num, col_num)] = value
            if idx == D.FIELD_INDEX.get("code") and url:
                value = f'=HYPERLINK("{_formula_string(url)}","{_formula_string(value)}")'

            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.font = row_font
            cell.alignment = CELL_ALIGN

        transfer_cell = ws.cell(row=row_num, column=transfer_col, value=None)
        transfer_cell.font = row_font
        transfer_cell.alignment = CELL_ALIGN

    _apply_week_frames(ws, week_keys, transfer_col)

    filter_start = ws.cell(row=TABLE_START_ROW, column=1).coordinate
    filter_end = ws.cell(
        row=max(TABLE_START_ROW, ws.max_row),
        column=transfer_col,
    ).coordinate
    ws.auto_filter.ref = f"{filter_start}:{filter_end}"

    ws.column_dimensions[get_column_letter(1)].width = WEEK_COLUMN_WIDTH

    for col_num, idx in enumerate(selected_indices, start=TABLE_START_COL):
        is_stations = idx == stations_idx
        width = STATION_COLUMN_WIDTH if is_stations else D.HEADERS[idx][1] + 2
        max_content_width = max(
            len(display_values.get((row_num, col_num), str(ws.cell(row=row_num, column=col_num).value or "")))
            for row_num in range(TABLE_START_ROW, ws.max_row + 1)
        )
        ws.column_dimensions[get_column_letter(col_num)].width = (
            width if is_stations else min(max(width, max_content_width + 2), 80)
        )

    ws.column_dimensions[get_column_letter(transfer_col)].width = TRANSFER_COLUMN_WIDTH

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
# ─── END OF sessions_to_xlsx_bytes() ─────────────────────────────────────────


def write_sessions_xlsx(
    rows: list[D.Row],
    pretty_print: str | list[str],
    operator_bindings: dict[str, str],
    operator_colors: dict[str, str],
    output_stream: BinaryIO,
) -> None:
    """
    Write visible session rows as XLSX to the provided binary output stream.
    """

    output_stream.write(
        sessions_to_xlsx_bytes(rows, pretty_print, operator_bindings, operator_colors)
    )
# ─── END OF write_sessions_xlsx() ────────────────────────────────────────────
