import io

from openpyxl import load_workbook

from ivs_sessions_browser.xlsx_export import _selected_stations, sessions_to_xlsx_bytes


def test_selected_stations_omits_removed_stations_and_keeps_spacing() -> None:
    assert _selected_stations("FtNnWzNs") == "Nn      Ns"
    assert _selected_stations("FtNnWz [NsHb]") == "Nn"
    assert _selected_stations("FtWz [NnNs]") == ""
    assert _selected_stations("FtNsWz [NnHb]") == "        Ns"
    assert _selected_stations("FtWzNs [Hb]") == "        Ns"
    assert _selected_stations("FtWz [Hb]") == ""


def test_xlsx_stations_column_contains_only_nn_and_ns() -> None:
    values = [
        "A", "R1", "R1234", "2026-10-27 12:00", "300", "24:00",
        "FtNnWz [NsHb]", "", "", "", "", "",
    ]
    rows = [(values, None, {})]

    workbook_bytes = sessions_to_xlsx_bytes(rows, ["STATIONS"], {}, {})
    workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=False)

    sheet = workbook["Sessions"]
    assert sheet["A5"].value == "Week"
    assert sheet["B5"].value == " Stations "
    assert sheet["C5"].value == "Transfer done"
    assert sheet["A6"].value == "Week 44"
    assert sheet["B6"].value == "Nn"
    assert sheet["C6"].value is None
    assert sheet.column_dimensions["A"].width == 11
    assert sheet.column_dimensions["B"].width == 10
    assert sheet.column_dimensions["C"].width == 14
    assert sheet["B5"].font.name == "Courier New"
    assert sheet["B6"].font.name == "Courier New"


def test_xlsx_merges_and_frames_each_week() -> None:
    def row(start: str) -> tuple[list[str], None, dict]:
        values = ["A", "R1", "R1234", start, "", "", "NnNs", "", "", "", "", ""]
        return values, None, {}

    rows = [
        row("2026-11-02 05:30"),
        row("2026-11-05 18:30"),
        row("2026-11-09 17:00"),
    ]
    workbook = load_workbook(
        io.BytesIO(sessions_to_xlsx_bytes(rows, ["OP", "STATIONS"], {}, {})),
        data_only=False,
    )
    sheet = workbook["Sessions"]

    assert "A6:A7" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    assert sheet["A6"].value == "Week 45"
    assert sheet["A8"].value == "Week 46"
    assert sheet["A6"].border.top.style == "medium"
    assert sheet["A7"].border.bottom.style == "medium"
    assert sheet["B6"].border.top.style == "medium"
    assert sheet["B7"].border.bottom.style == "medium"
    assert sheet["D6"].border.right.style == "medium"
    assert sheet["D7"].border.bottom.style == "medium"
